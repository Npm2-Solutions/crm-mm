# Copyright (c) 2026, NPM2 Solutions Srl and contributors
# For license information, please see license.txt

"""The food tables and the exercises dataset as the libraries keep them, without
a site (design.md, "I piani", "Le librerie"; ricerca-design.md §2.3).

- **A table is a sheet**: CIQUAL as ANSES publishes it (Excel, French or English),
  the Italian tables as the centre receives them with the licence (BDA-IEO) or the
  written permission (CREA), any other in the same shape: a row a food, a column a
  value for 100 g. The columns are recognised by their names; what is not
  recognised the person importing says.
- **The numbers are the table's**: "4,63" is 4.63; "-" is not known and is not
  counted; "traces" and "< 0,15" (under what the laboratory measures) count as
  nothing. Energy only in kJ becomes kcal; carbohydrates "by difference" lose the
  fibre they include. A value no food can have (more than 100 g in 100 g) is
  dropped. A row without a name or without a single value is not a food.
- **Energy the table leaves out** (CIQUAL does for a food in four, when one
  component is missing) is computed from proteins, carbohydrates, fats, fibre and
  alcohol with the factors of the Regulation EU 1169/2011 - the table's own way,
  within a few kcal where both exist - and the food says so.
- **The group is the centre's**: the table's own category becomes one of the
  library's groups by its words, in Italian, French or English; the person
  importing sees each category with the group it became and may change it.
- **exercises-dataset** (hasaneyldrm, MIT data; media © Gym visual with its
  written authorisation to NPM2 Solutions): the name in English, how it is done
  in Italian in steps, the body part, the equipment and the muscles in the
  centre's language; the animation and the picture only from where the agency
  hosts them. The assistant never touches them.
"""

from __future__ import annotations

import csv
import io
import math
import re
import unicodedata

from crm.clinica.piani_regole import GRUPPI

CIBO_CAMPI = (
	"code",
	"name",
	"group",
	"kcal",
	"kj",
	"protein_g",
	"carbs_g",
	"fat_g",
	"fibre_g",
	"alcohol_g",
)
#: What the library keeps of a food: its values for 100 g.
VALORI = ("kcal", "protein_g", "carbs_g", "fat_g", "fibre_g")
ALTRO = "Other"
#: The most a table's sheet may hold: a table is a thousand foods, not a database.
MAX_RIGHE = 20000
#: How many rows the header is looked for in: a title or a note may come first.
RIGHE_INTESTAZIONE = 15
KCAL_MAX = 1000
KJ_PER_KCAL = 4.184
#: Energy from the nutrients (Regulation EU 1169/2011, annex XIV), kcal per gram:
#: what a table computes itself, for the foods it leaves without.
FATTORI_UE = {"protein_g": 4, "carbs_g": 4, "fat_g": 9, "fibre_g": 2, "alcohol_g": 7}


# ------------------------------------------------------------------ words


def normalizza(testo) -> str:
	"""Lower case, without accents or signs: "Protéines, N x 6.25 (g/100 g)" is
	"proteines n x 6 25 g 100 g"."""
	testo = str(testo if testo is not None else "")
	for legatura, lettere in (("œ", "oe"), ("Œ", "oe"), ("æ", "ae"), ("Æ", "ae"), ("ß", "ss")):
		testo = testo.replace(legatura, lettere)
	testo = unicodedata.normalize("NFKD", testo).encode("ascii", "ignore").decode()
	return re.sub(r"[^a-z0-9]+", " ", testo.lower()).strip()


def _parole(testo: str, radici) -> bool:
	"""Whether a word of ``testo`` starts with one of ``radici``; a root ending in
	"$" is a whole word."""
	for radice in radici:
		intera = radice.endswith("$")
		radice = radice.rstrip("$")
		if re.search(r"(?<![a-z0-9])" + re.escape(radice) + (r"(?![a-z0-9])" if intera else ""), testo):
			return True
	return False


# ------------------------------------------------------------------ reading a sheet


def leggi_foglio(nome_file: str, contenuto: bytes) -> list[list]:
	"""The rows of a table's first sheet: Excel (.xlsx, .xls) or text (.csv, .txt)
	with commas, semicolons or tabs."""
	estensione = (nome_file or "").rsplit(".", 1)[-1].lower()
	if estensione == "xlsx" or contenuto[:2] == b"PK":
		from openpyxl import load_workbook

		libro = load_workbook(io.BytesIO(contenuto), read_only=True, data_only=True)
		foglio = libro.worksheets[0]
		righe = []
		for riga in foglio.iter_rows(values_only=True):
			righe.append(list(riga))
			if len(righe) > MAX_RIGHE:
				break
		libro.close()
		return righe
	if estensione == "xls" or contenuto[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
		import xlrd

		foglio = xlrd.open_workbook(file_contents=contenuto).sheet_by_index(0)
		return [foglio.row_values(i) for i in range(min(foglio.nrows, MAX_RIGHE + 1))]
	return _csv(contenuto)


def _csv(contenuto: bytes) -> list[list]:
	for codifica in ("utf-8-sig", "cp1252", "latin-1"):
		try:
			testo = contenuto.decode(codifica)
			break
		except UnicodeDecodeError:
			continue
	campione = testo[:20000]
	try:
		separatore = csv.Sniffer().sniff(campione, delimiters=";,\t|").delimiter
	except csv.Error:
		separatore = ";" if campione.count(";") > campione.count(",") else ","
	righe = []
	for riga in csv.reader(io.StringIO(testo), delimiter=separatore):
		righe.append(riga)
		if len(righe) > MAX_RIGHE:
			break
	return righe


# ------------------------------------------------------------------ the columns


def _gruppo_ciqual(testo: str) -> int | None:
	# CIQUAL's sub-group, then its group: the third level adds names, not groups
	for livello, peso in (("alim ssgrp nom", 0), ("alim grp nom", 1)):
		if testo.startswith(livello):
			return peso
	return None


def _priorita(campo: str, testo: str) -> int | None:
	"""How well a column's name says it is ``campo``: lower is better, None no."""
	parole = testo.split()
	prima = parole[0] if parole else ""
	if campo == "code":
		if testo in ("alim code", "codice", "codice alimento", "code", "food code", "cod", "id"):
			return 0
		if testo in ("fdc id", "ndb no", "ndb number", "codice bda", "codice crea", "id alimento"):
			return 1
		return None
	if campo == "name":
		if testo.startswith("alim nom") and not testo.startswith("alim nom sci"):
			return 0
		if testo in (
			"nome",
			"nome alimento",
			"alimento",
			"denominazione",
			"descrizione",
			"name",
			"food",
			"food name",
			"description",
			"food description",
			"nom",
			"aliment",
		):
			return 1
		return None
	if campo == "group":
		ciqual = _gruppo_ciqual(testo)
		if ciqual is not None:
			return ciqual
		if testo in ("sottogruppo", "sottocategoria", "subgroup", "sub group", "subcategory"):
			return 3
		if testo in (
			"categoria",
			"gruppo",
			"gruppo alimentare",
			"categoria alimentare",
			"categoria merceologica",
			"category",
			"group",
			"food group",
			"food category",
			"groupe",
			"categorie",
		):
			return 4
		return None
	if campo == "kcal":
		if "kcal" in parole:
			return 0 if "1169" in parole else 1
		if testo in ("calorie", "calories", "energia kcal"):
			return 2
		return None
	if campo == "kj":
		if "kj" in parole:
			return 0 if "1169" in parole else 1
		return None
	if campo == "protein_g":
		if prima not in ("protein", "proteins", "proteine", "proteines", "proteina", "proteinas"):
			return None
		if any(p in parole for p in ("animal", "animali", "animales", "vegetal", "vegetali", "vegetales")):
			return None
		return 1 if ("crude" in parole or "6 25" in testo) else 0
	if campo == "carbs_g":
		if prima not in ("carboidrati", "carbohydrate", "carbohydrates", "glucides", "carbohidratos"):
			return None
		if "difference" in parole:
			return 2
		if "totali" in parole or "total" in parole:
			return 1
		return 0
	if campo == "fat_g":
		if any(
			p in parole for p in ("acidi", "acides", "fatty", "saturi", "saturated", "animali", "vegetali")
		):
			return None
		if prima in ("lipidi", "lipides", "lipids", "fat", "fats", "grassi") or testo.startswith(
			"total lipid"
		):
			return 0
		return None
	if campo == "fibre_g":
		if prima not in ("fibra", "fibre", "fibres", "fiber", "fibers"):
			return None
		if any(p in parole for p in ("solubile", "insolubile", "soluble", "insoluble")):
			return None
		return 0
	if campo == "alcohol_g":
		return 0 if prima in ("alcool", "alcol", "alcohol", "ethanol", "etanolo") else None
	return None


def lingua_delle_colonne(intestazioni: list) -> str:
	"""The language of a table, from its columns: it, fr or en."""
	testi = [normalizza(t) for t in intestazioni]
	if any(t.startswith("alim nom fr") or t.startswith("alim grp nom fr") for t in testi):
		return "fr"
	if any(t.startswith("alim nom eng") or t.startswith("alim grp nom eng") for t in testi):
		return "en"
	parole = {p for t in testi for p in t.split()}
	if parole & {"proteine", "carboidrati", "lipidi", "fibra", "energia", "nome", "codice", "alimento"}:
		return "it"
	if parole & {"proteines", "glucides", "lipides", "fibres", "energie", "nom", "aliment"}:
		return "fr"
	return "en"


def riconosci(intestazioni: list) -> dict:
	"""Which column is which: ``{campo: indice}``; the group is a list of columns,
	the most precise first (CIQUAL's sub-group, then its group). A column serves
	one field only."""
	testi = [normalizza(t) for t in intestazioni]
	mappa: dict = {}
	usate: set[int] = set()
	for campo in CIBO_CAMPI:
		candidati = []
		for indice, testo in enumerate(testi):
			if indice in usate or not testo:
				continue
			priorita = _priorita(campo, testo)
			if priorita is not None:
				candidati.append((priorita, indice))
		if not candidati:
			continue
		candidati.sort()
		if campo == "group":
			mappa[campo] = [indice for _p, indice in candidati[:3]]
			usate.update(mappa[campo])
		else:
			mappa[campo] = candidati[0][1]
			usate.add(candidati[0][1])
	if "carbs_g" in mappa and "difference" in testi[mappa["carbs_g"]].split():
		# USDA's carbohydrates by difference hold the fibre: it is taken out
		mappa["carbs_with_fibre"] = True
	return mappa


def fonte_probabile(intestazioni: list) -> str | None:
	"""The source a table's columns tell: CIQUAL names its own."""
	if any(normalizza(t) == "alim code" for t in intestazioni):
		return "CIQUAL"
	return None


def trova_intestazione(righe: list[list]) -> int | None:
	"""The row with the columns' names: the first that names a food and a value."""
	for indice, riga in enumerate(righe[:RIGHE_INTESTAZIONE]):
		mappa = riconosci(riga)
		if "name" in mappa and any(campo in mappa for campo in ("kcal", "kj", "protein_g", "fat_g")):
			return indice
	return None


# ------------------------------------------------------------------ the numbers


_MANCA = {"", "-", "--", "nd", "n d", "na", "n a", "nr", "?", "x"}


def numero(valore) -> float | None:
	"""A value for 100 g as a table writes it: None when not known, 0 for traces
	and for what is under the laboratory's limit."""
	if valore is None or isinstance(valore, bool):
		return None
	if isinstance(valore, int | float):
		return float(valore) if math.isfinite(valore) and valore >= 0 else None
	testo = str(valore).strip().lower().replace(" ", " ")
	if normalizza(testo) in _MANCA:
		return None
	if testo.startswith("tr") or testo.startswith("<"):
		return 0.0
	testo = testo.replace(" ", "")
	if "," in testo and "." in testo:
		# the last sign is the decimal one: 1.234,5 and 1,234.5
		migliaia = "." if testo.rfind(",") > testo.rfind(".") else ","
		testo = testo.replace(migliaia, "")
	testo = testo.replace(",", ".")
	try:
		numero_letto = float(testo)
	except ValueError:
		return None
	return numero_letto if math.isfinite(numero_letto) and numero_letto >= 0 else None


def _cella(riga: list, indice) -> object:
	if indice is None or not isinstance(indice, int) or indice >= len(riga):
		return None
	return riga[indice]


def _testo(valore) -> str:
	if valore is None:
		return ""
	if isinstance(valore, float) and valore.is_integer():
		# a code Excel keeps as a number: 25601.0 is "25601"
		return str(int(valore))
	return str(valore).strip()


# ------------------------------------------------------------------ the groups

#: The library's group from a category's words, in order: the first that says it
#: wins. "Vegetable oils" are oils, "fruits de mer" fish, "fruit juices" drinks.
REGOLE_GRUPPI: tuple[tuple[str, tuple[str, ...], frozenset | None], ...] = (
	(
		ALTRO,
		(
			"baby",
			"infantil",
			"infanzia",
			"substitu",
			"sostitut",
			"ricett",
			"recipe",
			"piatt",
			"dish",
			"plats",
			"entree",
			"starter",
			"savoury pastr",
			"savory pastr",
			"feuillete",
			"sandwich",
			"pizza",
			"soup",
			"zupp",
			"prodotti vari",
			"varie$",
			"fast food",
		),
		None,
	),
	(
		"Oils and fats",
		(
			"oil",
			"oli$",
			"olio",
			"huile",
			"fats$",
			"fat$",
			"grassi",
			"matieres grasses",
			"butter",
			"burro",
			"beurre",
			"margarin",
			"strutto",
			"lard",
		),
		None,
	),
	(
		"Drinks",
		(
			"beverage",
			"drink",
			"bevand",
			"bibit",
			"boisson",
			"juice",
			"succh",
			"jus$",
			"water",
			"acqua",
			"acque",
			"eau$",
			"eaux",
			"wine",
			"vino",
			"vini$",
			"beer",
			"birr",
			"coffee",
			"caffe",
			"tea$",
			"tisan",
			"alcoholic",
			"alcolic",
			"alcoolis",
			"liquor",
		),
		None,
	),
	(
		"Cereals and tubers",
		("savoury biscuit", "savory biscuit", "biscuits aperitif", "cracker", "grissin", "fette biscottate"),
		None,
	),
	(
		"Sweets",
		(
			"sweet",
			"dolc",
			"sugar",
			"zuccher",
			"sucre",
			"chocolat",
			"cioccolat",
			"confection",
			"confiser",
			"honey",
			"miele",
			"miel$",
			"jam$",
			"marmellat",
			"confettur",
			"confiture",
			"ice cream",
			"gelat",
			"glace",
			"sorbet",
			"dessert",
			"cake",
			"torte",
			"gateau",
			"pastr",
			"pasticcer",
			"patisser",
			"viennois",
			"biscuit",
			"biscott",
			"merendin",
			"caramell",
			"cereal bar",
			"barres cereal",
			"barrett",
		),
		None,
	),
	(
		"Nuts and seeds",
		(
			"nuts$",
			"nut$",
			"seed",
			"frutta secca",
			"frutta a guscio",
			"semi$",
			"semi oleos",
			"oleagin",
			"fruits a coque",
			"graine",
			"noci$",
			"nocciol",
			"mandorl",
			"arachid",
			"pistacch",
			"peanut",
		),
		None,
	),
	(
		"Legumes",
		("legumi", "legumineuse", "pulse", "fagiol", "ceci$", "lenticch", "bean", "lentil"),
		None,
	),
	# English "legumes" are pulses; French "légumes" are vegetables
	("Legumes", ("legume",), frozenset({"it", "en"})),
	(
		"Fish",
		(
			"fish",
			"pesc",
			"seafood",
			"shellfish",
			"mollus",
			"crustace",
			"crostace",
			"frutti di mare",
			"fruits de mer",
			"poisson",
		),
		None,
	),
	("Eggs", ("egg", "uov", "uova", "oeuf"), None),
	(
		"Meat",
		(
			"meat",
			"carn",
			"salum",
			"pollam",
			"poultry",
			"insaccat",
			"sausage",
			"delicatessen",
			"frattagli",
			"offal",
			"viande",
			"volaille",
			"charcuter",
			"abats",
			"prosciutt",
			"ham$",
		),
		None,
	),
	(
		"Milk and dairy",
		(
			"milk",
			"latte$",
			"latticin",
			"lattier",
			"dairy",
			"chees",
			"formagg",
			"yogurt",
			"yoghurt",
			"lait",
			"fromage",
			"yaourt",
		),
		None,
	),
	(
		"Cereals and tubers",
		(
			"cereal",
			"cereale",
			"bread",
			"pane$",
			"pani$",
			"pain$",
			"pains$",
			"pasta",
			"pates",
			"rice",
			"riso",
			"riz$",
			"flour",
			"farin",
			"grain",
			"potato",
			"patat",
			"pomme de terre",
			"pommes de terre",
			"tuber",
			"prodotti da forno",
		),
		None,
	),
	(
		"Vegetables",
		("vegetabl", "verdur", "ortagg", "mushroom", "fungh", "champignon"),
		None,
	),
	("Vegetables", ("legumes",), frozenset({"fr"})),
	("Fruit", ("fruit", "frutt"), None),
)


def gruppo_da(testi, lingua: str = "it") -> str:
	"""The library's group of a food, from its table's categories, the most
	precise first: the first that says something decides; else "Other"."""
	for testo in testi:
		normale = normalizza(testo)
		if not normale:
			continue
		for gruppo, radici, lingue in REGOLE_GRUPPI:
			if lingue is not None and lingua not in lingue:
				continue
			if _parole(normale, radici):
				return gruppo
	return ALTRO


# ------------------------------------------------------------------ the foods


def energia(valori: dict, alcol: float | None = None) -> float | None:
	"""The kcal of 100 g from its nutrients, with the factors of the Regulation EU
	1169/2011: only when proteins, carbohydrates and fats are all known; fibre
	and alcohol count when they are."""
	if any(valori.get(campo) is None for campo in ("protein_g", "carbs_g", "fat_g")):
		return None
	totale = sum(
		FATTORI_UE[campo] * (valori.get(campo) or 0) for campo in ("protein_g", "carbs_g", "fat_g", "fibre_g")
	)
	totale += FATTORI_UE["alcohol_g"] * (alcol or 0)
	return round(totale, 1)


def _valori(riga: list, mappa: dict) -> tuple[dict[str, float | None], int, bool]:
	"""The values for 100 g of a row, how many were dropped as impossible, and
	whether the energy was computed from the nutrients."""
	valori = {campo: numero(_cella(riga, mappa.get(campo))) for campo in ("kcal", *VALORI[1:])}
	scartati = 0
	if valori["kcal"] is None and mappa.get("kj") is not None:
		kj = numero(_cella(riga, mappa["kj"]))
		if kj is not None:
			valori["kcal"] = round(kj / KJ_PER_KCAL, 1)
	if valori["kcal"] is not None and valori["kcal"] > KCAL_MAX:
		valori["kcal"] = None
		scartati += 1
	for campo in VALORI[1:]:
		if valori[campo] is not None and valori[campo] > 100:
			valori[campo] = None
			scartati += 1
	if mappa.get("carbs_with_fibre") and valori["carbs_g"] is not None and valori["fibre_g"] is not None:
		valori["carbs_g"] = round(max(valori["carbs_g"] - valori["fibre_g"], 0), 2)
	calcolata = False
	if valori["kcal"] is None:
		alcol = numero(_cella(riga, mappa.get("alcohol_g")))
		valori["kcal"] = energia(valori, alcol if alcol is not None and alcol <= 100 else None)
		calcolata = valori["kcal"] is not None
	return valori, scartati, calcolata


def categoria(riga: list, mappa: dict) -> list[str]:
	"""A row's categories in the table, the most precise first, without the empty
	ones ("-" is one)."""
	indici = mappa.get("group")
	if indici is None:
		return []
	if isinstance(indici, int):
		indici = [indici]
	testi = []
	for indice in indici:
		testo = _testo(_cella(riga, indice))
		if normalizza(testo) and testo not in testi:
			testi.append(testo)
	return testi


def alimenti(righe: list[list], mappa: dict, lingua: str = "it") -> dict:
	"""The foods of a table's rows (the header excluded): each with its code, name,
	category, group and values for 100 g. A row without a name or without a single
	value is left out, and counted; a code seen twice is taken once."""
	cibi: list[dict] = []
	senza_nome = senza_valori = doppi = scartati = 0
	visti: set[str] = set()
	for riga in righe:
		if not isinstance(riga, list | tuple):
			continue
		riga = list(riga)
		nome = re.sub(r"\s+", " ", _testo(_cella(riga, mappa.get("name"))))[:140]
		if not nome:
			if any(_testo(c) for c in riga):
				senza_nome += 1
			continue
		valori, fuori, calcolata = _valori(riga, mappa)
		scartati += fuori
		if all(valore is None for valore in valori.values()):
			senza_valori += 1
			continue
		codice = _testo(_cella(riga, mappa.get("code")))[:40] or None
		chiave = codice or normalizza(nome)
		if chiave in visti:
			doppi += 1
			continue
		visti.add(chiave)
		categorie = categoria(riga, mappa)
		cibi.append(
			{
				"code": codice,
				"name": nome,
				"category": (categorie[0] if categorie else "")[:140],
				"group": gruppo_da(categorie, lingua),
				**valori,
				"kcal_computed": calcolata,
			}
		)
	return {
		"foods": cibi,
		"without_name": senza_nome,
		"without_values": senza_valori,
		"twice": doppi,
		"dropped_values": scartati,
	}


def categorie(cibi: list[dict]) -> list[dict]:
	"""Each category of a table with the group it became and how many foods it
	holds, in the table's order: what the person importing checks."""
	viste: dict[str, dict] = {}
	for cibo in cibi:
		chiave = cibo.get("category") or ""
		if chiave not in viste:
			viste[chiave] = {"category": chiave, "group": cibo["group"], "count": 0}
		viste[chiave]["count"] += 1
	return list(viste.values())


def applica_gruppi(cibi: list[dict], scelte: dict | None) -> list[dict]:
	"""The groups the person importing chose for the categories; the rest as read."""
	scelte = {k: v for k, v in (scelte or {}).items() if v in GRUPPI}
	for cibo in cibi:
		scelto = scelte.get(cibo.get("category") or "")
		if scelto:
			cibo["group"] = scelto
	return cibi


# ------------------------------------------------------------------ exercises-dataset

DATASET = "exercises-dataset"
#: The dataset, at the version the import was written on (18/03/2026).
DATASET_URL = (
	"https://raw.githubusercontent.com/hasaneyldrm/exercises-dataset/"
	"7455efae41b330c265e7cd4b78dfa848e7ce5ebd/data/exercises.json"
)
PARTI_DATASET = {
	"back": "Back",
	"cardio": "Full body",
	"chest": "Chest",
	"lower arms": "Arms",
	"lower legs": "Legs",
	"neck": "Neck",
	"shoulders": "Shoulders",
	"upper arms": "Arms",
	"upper legs": "Legs",
	"waist": "Core",
}

#: The dataset's words for equipment and muscles, in Italian: a few dozen, read
#: once. A word not here stays in English.
ITALIANO = {
	# equipment
	"body weight": "corpo libero",
	"dumbbell": "manubri",
	"cable": "cavi",
	"barbell": "bilanciere",
	"leverage machine": "macchina a leva",
	"band": "elastico",
	"smith machine": "multipower",
	"kettlebell": "kettlebell",
	"weighted": "sovraccarico",
	"stability ball": "fitball",
	"ez barbell": "bilanciere EZ",
	"assisted": "assistito",
	"sled machine": "macchina a slitta",
	"medicine ball": "palla medica",
	"rope": "corda",
	"roller": "rullo",
	"resistance band": "banda elastica",
	"bosu ball": "bosu",
	"olympic barbell": "bilanciere olimpico",
	"wheel roller": "ruota per addominali",
	"upper body ergometer": "ergometro a braccia",
	"skierg machine": "SkiErg",
	"hammer": "martello",
	"stationary bike": "cyclette",
	"tire": "pneumatico",
	"trap bar": "trap bar",
	"elliptical machine": "ellittica",
	"stepmill machine": "stepmill",
	# muscles
	"abs": "addominali",
	"abdominals": "addominali",
	"lower abs": "addominali bassi",
	"pectorals": "pettorali",
	"chest": "petto",
	"upper chest": "parte alta del petto",
	"biceps": "bicipiti",
	"triceps": "tricipiti",
	"glutes": "glutei",
	"delts": "deltoidi",
	"deltoids": "deltoidi",
	"rear deltoids": "deltoidi posteriori",
	"shoulders": "spalle",
	"upper back": "parte alta della schiena",
	"back": "schiena",
	"lower back": "zona lombare",
	"lats": "dorsali",
	"latissimus dorsi": "gran dorsale",
	"calves": "polpacci",
	"soleus": "soleo",
	"quads": "quadricipiti",
	"quadriceps": "quadricipiti",
	"hamstrings": "femorali",
	"forearms": "avambracci",
	"cardiovascular system": "sistema cardiovascolare",
	"spine": "erettori spinali",
	"traps": "trapezi",
	"trapezius": "trapezio",
	"adductors": "adduttori",
	"abductors": "abduttori",
	"inner thighs": "interno coscia",
	"serratus anterior": "dentato anteriore",
	"levator scapulae": "elevatore della scapola",
	"core": "core",
	"hip flexors": "flessori dell'anca",
	"obliques": "obliqui",
	"rhomboids": "romboidi",
	"brachialis": "brachiale",
	"ankles": "caviglie",
	"ankle stabilizers": "stabilizzatori della caviglia",
	"feet": "piedi",
	"hands": "mani",
	"wrists": "polsi",
	"wrist flexors": "flessori del polso",
	"wrist extensors": "estensori del polso",
	"grip muscles": "muscoli della presa",
	"rotator cuff": "cuffia dei rotatori",
	"sternocleidomastoid": "sternocleidomastoideo",
	"groin": "inguine",
	"shins": "tibiali",
}


def parola(testo: str, lingua: str) -> str:
	testo = (testo or "").strip()
	if lingua == "it":
		return ITALIANO.get(testo.lower(), testo)
	return testo


def _passi(record: dict, lingua: str) -> str:
	for codice in dict.fromkeys((lingua, "en")):
		passi = ((record.get("instruction_steps") or {}).get(codice)) or []
		passi = [str(p).strip() for p in passi if str(p).strip()]
		if passi:
			return "\n".join(f"{numero_passo}. {passo}" for numero_passo, passo in enumerate(passi, 1))
		testo = str(((record.get("instructions") or {}).get(codice)) or "").strip()
		if testo:
			return testo
	return ""


_PERCORSO_IMMAGINE = re.compile(r"^images/[A-Za-z0-9._-]+\.(jpg|jpeg|png)$")
_PERCORSO_ANIMAZIONE = re.compile(r"^videos/[A-Za-z0-9._-]+\.gif$")


def indirizzo_media(base: str | None, percorso) -> str | None:
	"""Where the agency hosts a picture of the dataset: its address plus the
	dataset's own path. Nothing when the agency hosts none."""
	base = (base or "").strip()
	percorso = str(percorso or "").strip()
	if not base or not percorso:
		return None
	if not (_PERCORSO_IMMAGINE.match(percorso) or _PERCORSO_ANIMAZIONE.match(percorso)):
		return None
	if not (base.startswith("https://") or (base.startswith("/") and not base.startswith("//"))):
		return None
	return base.rstrip("/") + "/" + percorso


def _percorso(valore, forma: re.Pattern) -> str | None:
	percorso = str(valore or "").strip()
	return percorso if forma.match(percorso) else None


def esercizio(record, lingua: str = "it") -> dict | None:
	"""An exercise of the dataset as the library keeps it; None for a record that
	is not one. Its pictures are paths in the dataset: where they are served from
	is the agency's, and may change."""
	if not isinstance(record, dict):
		return None
	codice = str(record.get("id") or "").strip()
	nome = re.sub(r"\s+", " ", str(record.get("name") or "")).strip()
	if not codice or not nome:
		return None
	nome = (nome[0].upper() + nome[1:])[:140]
	principale = parola(str(record.get("target") or ""), lingua)
	secondari = []
	for muscolo in record.get("secondary_muscles") or []:
		tradotto = parola(str(muscolo), lingua)
		if tradotto and tradotto not in secondari and tradotto != principale:
			secondari.append(tradotto)
	return {
		"code": codice[:40],
		"name": nome,
		"name_in_source": nome,
		"body_part": PARTI_DATASET.get(str(record.get("body_part") or "").strip().lower(), "Other"),
		"equipment": parola(str(record.get("equipment") or ""), lingua)[:140] or None,
		"primary_muscles": principale or None,
		"secondary_muscles": ", ".join(secondari) or None,
		"instructions": _passi(record, lingua) or None,
		"media_path": _percorso(record.get("image"), _PERCORSO_IMMAGINE),
		"animation_path": _percorso(record.get("gif_url"), _PERCORSO_ANIMAZIONE),
		# the media's owner, cited wherever they are shown
		"attribution": str(record.get("attribution") or "").strip()[:140] or None,
	}
