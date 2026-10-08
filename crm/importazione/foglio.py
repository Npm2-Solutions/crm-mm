# Copyright (c) 2026, NPM2 Solutions Srl and contributors
# For license information, please see license.txt

"""A sheet as another program saves it, without a site: Excel (.xlsx, .xls) or
text (.csv, .txt) with commas, semicolons or tabs, in the encodings Italian
programs write. The CRM's people and the clinic's food tables read it so."""

from __future__ import annotations

import csv
import io


def leggi(nome_file: str, contenuto: bytes, massimo: int = 20000) -> list[list]:
	"""The rows of the first sheet, at most ``massimo`` past the first."""
	estensione = (nome_file or "").rsplit(".", 1)[-1].lower()
	if estensione == "xlsx" or contenuto[:2] == b"PK":
		from openpyxl import load_workbook

		libro = load_workbook(io.BytesIO(contenuto), read_only=True, data_only=True)
		foglio = libro.worksheets[0]
		righe = []
		for riga in foglio.iter_rows(values_only=True):
			righe.append(list(riga))
			if len(righe) > massimo:
				break
		libro.close()
		return righe
	if estensione == "xls" or contenuto[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
		import xlrd

		foglio = xlrd.open_workbook(file_contents=contenuto).sheet_by_index(0)
		return [foglio.row_values(i) for i in range(min(foglio.nrows, massimo + 1))]
	return _csv(contenuto, massimo)


def _csv(contenuto: bytes, massimo: int) -> list[list]:
	for codifica in ("utf-8-sig", "cp1252", "latin-1"):
		try:
			testo = contenuto.decode(codifica)
			break
		except UnicodeDecodeError:
			continue
	campione = testo[:20000]
	try:
		separatore = csv.Sniffer().sniff(campione, delimiters=";,\t|").delimiter
	except csv.Error:
		separatore = ";" if campione.count(";") > campione.count(",") else ","
	righe = []
	for riga in csv.reader(io.StringIO(testo), delimiter=separatore):
		righe.append(riga)
		if len(righe) > massimo:
			break
	return righe
